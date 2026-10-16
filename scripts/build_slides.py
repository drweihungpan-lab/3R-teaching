#!/usr/bin/env python3
"""Build and check editable course slides. See materials/slides/README.md."""

import argparse
from dataclasses import dataclass
import os
from pathlib import Path
import re
import unicodedata
from urllib.parse import unquote, urlsplit
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


WIDTH, HEIGHT = 13.333333, 7.5
LEFT, BODY_TOP, BODY_BOTTOM = 0.65, 1.9, 6.93
BODY_WIDTH = WIDTH - 2 * LEFT
TITLE_SIZE, BODY_SIZE, CODE_SIZE = 32, 21, 18
INK, MUTED, ACCENT = "20344A", "53677A", "007F82"
INLINE = re.compile(r"`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|\*([^*]+)\*")


@dataclass(frozen=True)
class Span:
    text: str
    code: bool = False
    bold: bool = False
    italic: bool = False
    link: str | None = None


@dataclass
class Element:
    key: str
    spans: list[Span]
    code: bool = False
    bullet: bool = False

    @property
    def text(self):
        return "".join(span.text for span in self.spans)


@dataclass
class Section:
    title: Element
    elements: list[Element]


def inline(text):
    """Convert Markdown inline syntax to editable runs, retaining link targets."""
    spans, start = [], 0
    for match in INLINE.finditer(text):
        if match.start() > start:
            spans.append(Span(text[start:match.start()]))
        code, label, target, bold, italic = match.groups()
        spans.append(Span(code or label or bold or italic or "", code=code is not None,
                          bold=bold is not None, italic=italic is not None, link=target))
        start = match.end()
    if start < len(text):
        spans.append(Span(text[start:]))
    return spans or [Span("")]


def parse_source(source):
    """Read the course slide dialect; do not treat --- inside a fence as a slide."""
    sections, current, fence = [], None, None
    expect_title = True
    for number, line in enumerate(source.read_text(encoding="utf-8").splitlines(), 1):
        key = f"line-{number}"
        if fence is not None:
            if line.strip() == fence:
                fence = None
            else:
                current.elements.append(Element(key, [Span(line, code=True)], code=True))
            continue
        if line.startswith("```"):
            if current is None or expect_title:
                raise ValueError(f"Line {number}: code before a slide title")
            fence = "```"
            continue
        if line.strip() == "---":
            if current is None or expect_title:
                raise ValueError(f"Line {number}: empty slide")
            expect_title = True
            continue
        if not line.strip():
            continue
        heading = re.fullmatch(r"(#{1,2}) (.+)", line)
        if expect_title:
            if not heading:
                raise ValueError(f"Line {number}: expected # or ## slide title")
            current = Section(Element(key, inline(heading[2])), [])
            sections.append(current)
            expect_title = False
        elif heading:
            raise ValueError(f"Line {number}: add --- before the next slide title")
        else:
            if re.match(r"\s*(?:[>#]|\||[-*+]\s*\[[ xX]\])", line):
                raise ValueError(f"Line {number}: unsupported block syntax; see slide README")
            bullet = line.startswith(("- ", "* ", "+ "))
            current.elements.append(Element(key, inline(line[2:] if bullet else line),
                                            bullet=bullet))
    if fence is not None:
        raise ValueError("Unclosed code fence")
    if not sections or expect_title:
        raise ValueError("Source is empty or ends with an empty slide")
    if any(not section.elements or not any(e.text.strip() for e in section.elements)
           for section in sections):
        raise ValueError("Every source slide needs a title and nonempty body")
    return sections


def link_address(target, source, output):
    parts = urlsplit(target)
    if parts.scheme or parts.netloc or target.startswith("#"):
        return target
    local = source.parent / unquote(parts.path)
    if not local.exists():
        raise ValueError(f"Missing source link: {target}")
    relative = Path(os.path.relpath(local.resolve(), output.parent.resolve())).as_posix()
    return relative + (f"#{parts.fragment}" if parts.fragment else "")


def character_width(char, size, code=False):
    """Conservative point estimates; these are layout guards, not font rendering."""
    if unicodedata.east_asian_width(char) in ("W", "F"):
        return size * 1.1
    if code:
        return size * 0.64
    if char in "MW@%mw":
        return size * 1.04
    if char in " il.,:;!|'`":
        return size * 0.36
    return size * (0.8 if char.isupper() else 0.65)


def wrapped_ranges(element, size, width):
    """Return exact source substrings; soft breaks never remove whitespace/tokens."""
    text = element.text
    if not text:
        return [(0, 0)]
    widths = [character_width(c, size, element.code or span.code)
              for span in element.spans for c in span.text]
    ranges, start = [], 0
    while start < len(text):
        end, used, last_space = start, 0, None
        while end < len(text) and used + widths[end] <= width:
            used += widths[end]
            if text[end].isspace():
                last_space = end + 1
            end += 1
        if end == start:
            raise ValueError("A character is wider than the available text area")
        if end < len(text) and last_space is not None and last_space > start:
            end = last_space
        ranges.append((start, end))
        start = end
    return ranges


def color(value):
    return RGBColor.from_string(value)


def textbox(slide, name, x, y, width, height, size, fill=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill)
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False  # All wrapping is explicit and included in the height budget.
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Pt(7)
    frame.margin_top = frame.margin_bottom = Pt(3)
    paragraph = frame.paragraphs[0]
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color(INK)
    paragraph.line_spacing = Pt(size * 1.32)
    paragraph.space_before = paragraph.space_after = Pt(0)
    return shape, paragraph


def add_runs(paragraph, element, ranges, size, source, output):
    for row, (start, end) in enumerate(ranges):
        if row:
            paragraph.add_line_break()  # Soft breaks can be removed to recover exact text.
        offset = 0
        for span in element.spans:
            left, right = max(start, offset), min(end, offset + len(span.text))
            if right > left:
                run = paragraph.add_run()
                run.text = span.text[left - offset:right - offset]
                run.font.name = "Courier New" if element.code or span.code else "Arial"
                run.font.size = Pt(size)
                run.font.bold = span.bold
                run.font.italic = span.italic
                run.font.color.rgb = color(ACCENT if span.link else INK)
                if span.link:
                    run.hyperlink.address = link_address(span.link, source, output)
            offset += len(span.text)


def add_page(prs, section, index, total, part, source, output):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color("FAFCFD")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.07))
    bar.name = "accent"
    bar.fill.solid()
    bar.fill.fore_color.rgb = color(ACCENT)
    bar.line.fill.background()
    _, label = textbox(slide, "section", LEFT, 0.22, BODY_WIDTH, 0.28, 11)
    label.text = f"SOURCE {index} / {total}" + (f"  ·  CONTINUED {part}" if part > 1 else "")
    label.font.color.rgb = color(MUTED)
    title_lines = wrapped_ranges(section.title, TITLE_SIZE, (BODY_WIDTH * 72 - 14) * 0.94)
    if len(title_lines) > 2:
        raise ValueError(f"Title needs more than two lines: {section.title.text}")
    _, paragraph = textbox(slide, f"title:{section.title.key}", LEFT, 0.58,
                           BODY_WIDTH, 1.27, TITLE_SIZE)
    paragraph.line_spacing = Pt(40)
    add_runs(paragraph, section.title, title_lines, TITLE_SIZE, source, output)
    _, footer = textbox(slide, "footer", LEFT, 7.12, BODY_WIDTH, 0.24, 10)
    footer.text = f"{source.name}   ·   {len(prs.slides)}"
    footer.font.color.rgb = color(MUTED)
    return slide


def build(sections, source, output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    prs.core_properties.title = sections[0].title.text
    prs.core_properties.subject = f"Editable slides generated from {source.name}"
    prs.core_properties.author = "Three-R course"
    prs.core_properties.keywords = "editable, Markdown, course"
    for index, section in enumerate(sections, 1):
        part, y = 1, BODY_TOP
        slide = add_page(prs, section, index, len(sections), part, source, output)
        for element in section.elements:
            size = CODE_SIZE if element.code else BODY_SIZE
            line_height = size * 1.32
            # Inset bullets inside their text box, rather than using negative geometry.
            inset = 22 if element.bullet else 0
            rows = wrapped_ranges(element, size, (BODY_WIDTH * 72 - 14 - inset) * 0.94)
            cursor = 0
            while cursor < len(rows):
                available = (BODY_BOTTOM - y) * 72
                capacity = int((available - 6) // line_height)
                if capacity < 1:
                    part += 1
                    slide = add_page(prs, section, index, len(sections), part, source, output)
                    y = BODY_TOP
                    continue
                chunk = rows[cursor:cursor + capacity]
                height = (len(chunk) * line_height + 6) / 72
                shape, paragraph = textbox(slide, f"content:{element.key}:{cursor}", LEFT, y,
                                            BODY_WIDTH, height, size,
                                            "EDF3F6" if element.code else None)
                if element.bullet:
                    properties = paragraph._p.get_or_add_pPr()
                    properties.set("marL", str(Pt(22)))
                    properties.set("indent", str(Pt(-17)))
                    bullet = OxmlElement("a:buChar")
                    bullet.set("char", "•" if cursor == 0 else " ")
                    properties.append(bullet)
                add_runs(paragraph, element, chunk, size, source, output)
                y += height + (0 if element.code else 0.07)
                cursor += len(chunk)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def validate(sections, source, output):
    """Check the saved artifact, including coverage/order, native text and slide bounds."""
    prs = Presentation(output)
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.0001:
        raise ValueError("Expected 16:9 slide dimensions")
    expected = [element for section in sections for element in section.elements]
    expected_titles = [section.title for section in sections]
    actual, order, titles, links = {}, [], [], set()
    body_shapes = 0
    for number, slide in enumerate(prs.slides, 1):
        has_body, title_count, previous_bottom = False, 0, Inches(BODY_TOP)
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or shape.width <= 0 or shape.height <= 0
                    or shape.left + shape.width > prs.slide_width + 2
                    or shape.top + shape.height > prs.slide_height + 2):
                raise ValueError(f"Slide {number}: out-of-bounds shape {shape.name}")
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                raise ValueError(f"Slide {number}: raster image found instead of editable content")
            if not shape.has_text_frame:
                continue
            text = shape.text.replace("\v", "")
            if shape.name.startswith("title:"):
                key = shape.name.split(":")[1]
                if not titles or titles[-1][0] != key:
                    titles.append((key, text))
                elif titles[-1][1] != text:
                    raise ValueError(f"Slide {number}: changed continuation title")
                title_count += 1
            if shape.name.startswith("content:"):
                key = shape.name.split(":")[1]
                if shape.top < previous_bottom - 2:
                    raise ValueError(f"Slide {number}: overlapping body shapes")
                if shape.top + shape.height > Inches(BODY_BOTTOM) + 2:
                    raise ValueError(f"Slide {number}: body exceeds content region")
                previous_bottom = shape.top + shape.height
                if not order or order[-1] != key:
                    order.append(key)
                actual[key] = actual.get(key, "") + text
                has_body = has_body or bool(text.strip())
                body_shapes += 1
                # Explicit row heights and width estimates guard against internal clipping.
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.font.size is None:
                        raise ValueError(f"Slide {number}: missing explicit body font size")
                    size = paragraph.font.size.pt
                    rows = paragraph.text.split("\v")
                    if len(rows) * size * 1.32 + 6 > shape.height / 12700 + 0.01:
                        raise ValueError(f"Slide {number}: text height exceeds its box")
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.hyperlink.address:
                        links.add(run.hyperlink.address)
        if not has_body or title_count != 1:
            raise ValueError(f"Slide {number}: empty/missing body or title")
    if titles != [(element.key, element.text) for element in expected_titles]:
        raise ValueError("Source titles are missing, changed, or out of sequence")
    if order != [element.key for element in expected]:
        raise ValueError("Source body elements are missing, duplicated, or out of sequence")
    for element in expected:
        if actual.get(element.key) != element.text:
            raise ValueError(f"Source text mismatch at {element.key}")
    expected_links = {link_address(span.link, source, output)
                      for section in sections for element in [section.title, *section.elements]
                      for span in element.spans if span.link}
    if links != expected_links:
        raise ValueError("Hyperlink targets differ from the source")
    with ZipFile(output) as archive:
        if archive.testzip():
            raise ValueError("Corrupt PPTX ZIP member")
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                etree.fromstring(archive.read(name))
        slide_parts = [name for name in archive.namelist()
                       if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        if len(slide_parts) != len(prs.slides):
            raise ValueError("Missing or orphaned slide XML")
    return (f"PASS: {len(sections)} source slides -> {len(prs.slides)} editable slides; "
            f"{len(expected)} source elements, {body_shapes} text boxes, {len(links)} links; "
            "exact text/order, shape bounds, body overlap/height budgets, XML/ZIP checked. "
            "No rendered inspection; font substitution and visual clipping need PowerPoint review.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True, type=Path, help="Markdown slide source")
    parser.add_argument("--output", required=True, type=Path, help="PowerPoint destination")
    parser.add_argument("--check-only", action="store_true", help="Validate an existing deck without writing")
    args = parser.parse_args()
    try:
        if args.source.resolve() == args.output.resolve():
            raise ValueError("Source and output must be different files")
        if args.output.suffix.lower() != ".pptx":
            raise ValueError("Output must have a .pptx extension")
        sections = parse_source(args.source)
        if not args.check_only:
            build(sections, args.source, args.output)
        print(validate(sections, args.source, args.output))
    except (OSError, ValueError, etree.XMLSyntaxError) as error:
        parser.exit(1, f"Slide build/check failed: {error}\n")


if __name__ == "__main__":
    main()
